from __future__ import annotations

import os
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .models import ConflictDetail, RequestRecord, RunSummary
from .normalize import safe_text

HEADER = PatternFill("solid", fgColor="1F4E78"); WHITE = Font(color="FFFFFF", bold=True)
GREEN = PatternFill("solid", fgColor="E2F0D9"); RED = PatternFill("solid", fgColor="FCE4D6"); YELLOW = PatternFill("solid", fgColor="FFF2CC")
RESULT_TYPE = "Tipo da atividade em conflito"; RESULT_MESSAGE = "Resultado da verificação"
TRACE_HEADERS = ["Arquivo de origem", "Caminho relativo", "Aba de origem", "Linha de origem", "Pasta de origem", "Perfil utilizado", "Identificador do registro", "Identificador da execução"]


def _write_sheet(workbook: Workbook, title: str, records: list[RequestRecord]) -> None:
    ws = workbook.create_sheet(title)
    data_headers = []
    for record in records:
        for header in record.values:
            if header not in data_headers: data_headers.append(header)
    headers = data_headers + [RESULT_TYPE, RESULT_MESSAGE] + TRACE_HEADERS
    ws.append(headers)
    for cell in ws[1]: cell.fill = HEADER; cell.font = WHITE
    for record in records:
        row = [safe_text(record.values.get(h, "")) for h in data_headers] + [record.result_type, record.result_message, record.source_file, record.relative_path, record.source_sheet, record.source_row, record.source_folder, record.profile_id, record.record_id, record.execution_id]
        ws.append(row)
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    for row in range(2, ws.max_row + 1):
        ws.cell(row, headers.index(RESULT_TYPE) + 1).alignment = Alignment(wrap_text=True, vertical="top")
        text = str(ws.cell(row, headers.index(RESULT_MESSAGE) + 1).value or "")
        ws.cell(row, headers.index(RESULT_MESSAGE) + 1).alignment = Alignment(wrap_text=True, vertical="top")
        if text.startswith("Conflito"): fill = RED
        elif text.startswith("Sem conflito"): fill = GREEN
        else: fill = YELLOW
        for cell in ws[row]: cell.fill = fill
    for index, header in enumerate(headers, 1):
        lengths = [len(str(header))] + [len(str(ws.cell(row, index).value or "")) for row in range(2, min(ws.max_row, 100) + 1)]
        length = max(lengths)
        ws.column_dimensions[get_column_letter(index)].width = min(max(length + 2, 12), 55)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, datetime): cell.number_format = "DD/MM/YYYY HH:MM"
    if ws.max_row > 1:
        table = Table(displayName=f"Tabela{title.replace(' ', '').replace('ã','a')}", ref=ws.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)


def _summary_sheet(workbook: Workbook, summary: RunSummary) -> None:
    ws = workbook.create_sheet("Resumo")
    rows = [("Identificador da execução", summary.execution_id), ("Início", summary.started_at), ("Fonte utilizada", summary.source_label), ("Status", summary.status), ("Arquivos encontrados", summary.files_found), ("Arquivos processados", summary.files_processed), ("Arquivos com erro", summary.files_with_error), ("Solicitações letivas", summary.letivas), ("Solicitações não letivas", summary.nao_letivas), ("Eventos com conflito", summary.conflicts), ("Eventos sem conflito", summary.no_conflicts), ("Eventos não verificados", summary.not_verified), ("Ocorrências de conflito", summary.conflict_occurrences)]
    for row in rows: ws.append(row)
    ws["A1"].fill = HEADER; ws["A1"].font = WHITE; ws.column_dimensions["A"].width = 32; ws.column_dimensions["B"].width = 80
    if summary.notes:
        ws.append([]); ws.append(["Pendências e avisos"]); ws.cell(ws.max_row, 1).fill = HEADER; ws.cell(ws.max_row, 1).font = WHITE
        for note in summary.notes: ws.append([note])


def _simple_table(workbook: Workbook, title: str, headers: list[str], rows: list[list[object]]) -> None:
    ws = workbook.create_sheet(title); ws.append(headers)
    for cell in ws[1]: cell.fill = HEADER; cell.font = WHITE
    for row in rows: ws.append(row)
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    for index, header in enumerate(headers, 1): ws.column_dimensions[get_column_letter(index)].width = min(max(len(header) + 2, 16), 55)


def export_report(output_dir: Path, summary: RunSummary, records: list[RequestRecord], details: list[ConflictDetail], file_rows: list[list[object]], issues: list[list[object]]) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    wb = Workbook(); wb.remove(wb.active)
    _summary_sheet(wb, summary)
    _write_sheet(wb, "Letivas", [r for r in records if r.request_type == "letiva"])
    _write_sheet(wb, "Não Letivas", [r for r in records if r.request_type == "nao_letiva"])
    _simple_table(wb, "Pendências", ["Categoria", "Origem", "Mensagem"], issues)
    _simple_table(wb, "Arquivos processados", ["Tipo", "Pasta", "Arquivo", "Status", "Mensagem"], file_rows)
    if details:
        _simple_table(wb, "Detalhamento de conflitos", ["ID do registro", "Tipo da solicitação", "Matrícula", "Data", "Início Evento", "Fim Evento", "Início DW", "Fim DW", "Tipo DW", "Atividade DW"], [[d.request_id, d.request_type, d.matricula, d.activity_date, d.event_start, d.event_end, d.dw_start, d.dw_end, d.dw_type, d.dw_description] for d in details])
    filename = f"consolidado_eventos_avulsos_{datetime.now():%Y%m%d_%H%M%S}_{summary.execution_id[:8]}.xlsx"
    target = output_dir / filename; temporary = target.with_suffix(".tmp.xlsx")
    wb.save(temporary)
    load_workbook(temporary, read_only=True).close()
    os.replace(temporary, target)
    return target
