from __future__ import annotations

from openpyxl.utils import column_index_from_string

from .models import Profile
from .normalize import normalized_header

REQUIRED = ("matricula", "data", "inicio", "fim")


def resolve_profile_sheet(workbook, profile: Profile) -> tuple[object, dict[str, int]]:
    candidates = [workbook[name] for name in profile.sheet_names if name in workbook.sheetnames] if profile.sheet_names else list(workbook.worksheets)
    matches = []
    for sheet in candidates:
        header_cells = next(sheet.iter_rows(min_row=profile.header_row, max_row=profile.header_row), ())
        headers = {normalized_header(cell.value): cell.column for cell in header_cells if cell.value is not None}
        resolved = {}
        valid = True
        for field in REQUIRED:
            spec = profile.fields.get(field, {})
            aliases = [normalized_header(a) for a in spec.get("aliases", [])]
            alias_cols = {headers[a] for a in aliases if a in headers}
            if len(alias_cols) > 1:
                valid = False; break
            if alias_cols:
                resolved[field] = alias_cols.pop()
            elif spec.get("column"):
                resolved[field] = column_index_from_string(spec["column"])
                if resolved[field] > sheet.max_column:
                    valid = False; break
            else:
                valid = False; break
        if valid:
            matches.append((sheet, resolved))
    if not matches:
        raise ValueError(f"Perfil '{profile.name}' não reconheceu nenhuma aba do arquivo.")
    if len(matches) > 1:
        names = ", ".join(item[0].title for item in matches)
        raise ValueError(f"Perfil '{profile.name}' é ambíguo nas abas: {names}.")
    return matches[0]
