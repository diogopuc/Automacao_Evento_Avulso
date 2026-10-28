from __future__ import annotations

from pathlib import Path
from uuid import uuid4

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .models import CancellationToken, Profile, RequestRecord
from .normalize import normalize_id, parse_date, parse_time
from .profiles import resolve_profile_sheet


def read_requests(path: Path, folder_name: str, request_type: str, profile: Profile, root: Path, execution_id: str, token: CancellationToken) -> tuple[list[RequestRecord], list[str]]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=False, keep_vba=path.suffix.lower() == ".xlsm")
    except Exception as exc:
        raise ValueError(f"Não foi possível abrir '{path.name}': {exc}") from exc
    sheet, fields = resolve_profile_sheet(workbook, profile)
    header_cells = next(sheet.iter_rows(min_row=profile.header_row, max_row=profile.header_row), ())
    headers = [str(cell.value).strip() if cell.value not in (None, "") else f"Coluna {cell.column}" for cell in header_cells]
    # Duplicate headings are made stable but never discarded.
    seen: dict[str, int] = {}
    safe_headers = []
    for header in headers:
        seen[header] = seen.get(header, 0) + 1
        safe_headers.append(header if seen[header] == 1 else f"{header} ({seen[header]})")
    records, warnings = [], []
    for row_number, row in enumerate(sheet.iter_rows(min_row=profile.first_data_row, values_only=False), start=profile.first_data_row):
        token.raise_if_cancelled()
        raw = [cell.value for cell in row]
        if not any(value not in (None, "") for value in raw):
            continue
        values = {safe_headers[index]: value for index, value in enumerate(raw) if index < len(safe_headers)}
        def cell_value(field):
            if field not in fields:
                return None
            index = fields[field] - 1
            return raw[index] if index < len(raw) else None
        record = RequestRecord(values=values, request_type=request_type, matricula=normalize_id(cell_value("matricula")), activity_date=parse_date(cell_value("data")), start_time=parse_time(cell_value("inicio")), end_time=parse_time(cell_value("fim")), source_file=path.name, relative_path=str(path.relative_to(root)), source_sheet=sheet.title, source_row=row_number, source_folder=folder_name, profile_id=profile.id, record_id=str(uuid4()), execution_id=execution_id)
        if not record.matricula: record.issues.append("Matrícula ausente ou inválida")
        if not record.activity_date: record.issues.append("Data ausente ou inválida")
        if not record.start_time or not record.end_time: record.issues.append("Horário de início ou término ausente ou inválido")
        elif record.end_time <= record.start_time: record.issues.append("Término igual ou anterior ao início")
        records.append(record)
    if not records:
        warnings.append("Nenhum registro de dados encontrado")
    return records, warnings


def preview_mapping(path: Path, profile: Profile, limit: int = 5) -> tuple[str, dict[str, str], list[list[object]]]:
    """Reads only a small preview for the explicit user review, without writing the source."""
    workbook = load_workbook(path, read_only=True, data_only=False, keep_vba=path.suffix.lower() == ".xlsm")
    sheet, fields = resolve_profile_sheet(workbook, profile)
    header_cells = next(sheet.iter_rows(min_row=profile.header_row, max_row=profile.header_row), ())
    header_values = {cell.column: cell.value for cell in header_cells}
    mapping = {field: f"{get_column_letter(column)} — {header_values.get(column, '')}" for field, column in fields.items()}
    rows = []
    for row in sheet.iter_rows(min_row=profile.first_data_row, max_row=profile.first_data_row + limit - 1, values_only=True):
        rows.append(list(row))
    return sheet.title, mapping, rows
