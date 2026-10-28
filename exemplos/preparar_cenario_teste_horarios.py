"""Preenche W/X nos arquivos principais para testar a automação.

Uso: python exemplos/preparar_cenario_teste_horarios.py
Altera somente 1. LETIVA e 2. NAO LETIVA, por solicitação explícita.
"""
from __future__ import annotations

import os
import sys
from datetime import time
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

SOURCE_FOLDERS = (ROOT / "1. LETIVA", ROOT / "2. NAO LETIVA")
def main() -> None:
    files = [path for folder in SOURCE_FOLDERS for path in folder.iterdir() if path.is_file() and not path.name.startswith("~$") and path.suffix.lower() in {".xlsx", ".xlsm"}]
    updated = 0
    for path in files:
        workbook = load_workbook(path, keep_vba=path.suffix.lower() == ".xlsm")
        try:
            sheet = workbook.active
            sheet.cell(1, 23).value = "Horário de início"
            sheet.cell(1, 24).value = "Horário de término"
            for row_number in range(2, sheet.max_row + 1):
                source_values = [sheet.cell(row_number, column).value for column in range(1, 23)]
                if not any(value not in (None, "") for value in source_values):
                    sheet.cell(row_number, 23).value = None
                    sheet.cell(row_number, 24).value = None
                    continue
                start, end = time(8, 0), time(9, 0)
                sheet.cell(row_number, 23).value = start
                sheet.cell(row_number, 24).value = end
                sheet.cell(row_number, 23).number_format = "HH:MM"
                sheet.cell(row_number, 24).number_format = "HH:MM"
                updated += 1
            temporary = path.with_name(f".{path.stem}.cenario_teste{path.suffix}")
            workbook.save(temporary)
            os.replace(temporary, path)
        finally:
            workbook.close()
    print(f"Linhas atualizadas: {updated}. Horário de teste aplicado: 08:00–09:00.")


if __name__ == "__main__":
    main()
