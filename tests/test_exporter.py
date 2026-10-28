from datetime import date, time
from zipfile import ZipFile

from openpyxl import load_workbook

from app.exporter import export_report
from app.models import RequestRecord, RunSummary

def test_export_reconciles_summary_and_keeps_traceability(tmp_path):
    record = RequestRecord({"Matrícula": "000123", "Data": date(2026,10,28)}, "letiva", "000123", date(2026,10,28), time(8), time(9), "origem.xlsx", "1. LETIVA/origem.xlsx", "Solicitações", 2, "1. LETIVA", "novo", "id-1", "exec-1", result_message="Sem conflito.")
    summary = RunSummary("exec-1", "2026-10-28T08:00:00", "Somente consolidação", letivas=1, no_conflicts=1)
    output = export_report(tmp_path, summary, [record], [], [["letiva", "1. LETIVA", "origem.xlsx", "Processado", ""]], [])
    wb = load_workbook(output, data_only=True)
    assert wb.sheetnames[:3] == ["Resumo", "Letivas", "Não Letivas"]
    ws = wb["Letivas"]
    headers = [cell.value for cell in ws[1]]
    assert ws.max_row == 2 and "Identificador do registro" not in headers and ws.cell(2, headers.index("Matrícula") + 1).value == "000123"
    with ZipFile(output) as archive:
        worksheet_xml = archive.read("xl/worksheets/sheet2.xml").decode("utf-8")
        assert "<autoFilter" not in worksheet_xml
