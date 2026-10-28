from datetime import date, time
from pathlib import Path

from openpyxl import Workbook

from app.models import CancellationToken, Profile
from app.reader import read_requests


def profile():
    return Profile("p", "Teste", ["Solicitações"], 1, 2, ["letiva"], {
        "matricula": {"column": "B", "aliases": ["matrícula"]},
        "data": {"column": "F", "aliases": ["data"]},
        "inicio": {"column": "W", "aliases": ["início"]},
        "fim": {"column": "X", "aliases": ["término"]},
    })

def test_reader_uses_configured_columns_and_keeps_zeroes(tmp_path):
    file = tmp_path / "solicitacoes.xlsx"; wb = Workbook(); ws = wb.active; ws.title = "Solicitações"
    headers = ["" for _ in range(24)]; headers[1], headers[5], headers[22], headers[23] = "Matrícula", "Data", "Início", "Término"; ws.append(headers)
    row = ["" for _ in range(24)]; row[1], row[5], row[22], row[23] = "000123", date(2026,10,28), time(8), time(9); ws.append(row); wb.save(file)
    records, warnings = read_requests(file, "1. LETIVA", "letiva", profile(), tmp_path, "exec", CancellationToken())
    assert not warnings and records[0].matricula == "000123" and not records[0].issues

def test_reader_rejects_ambiguous_template(tmp_path):
    file = tmp_path / "solicitacoes.xlsx"; wb = Workbook(); ws = wb.active; ws.title = "Solicitações"; ws.append(["Matrícula", "Data", "Início", "Término"]); wb.create_sheet("Outra").append(["Matrícula", "Data", "Início", "Término"]); wb.save(file)
    ambiguous_profile = profile(); ambiguous_profile.sheet_names = []
    try:
        read_requests(file, "1. LETIVA", "letiva", ambiguous_profile, tmp_path, "exec", CancellationToken())
    except ValueError as exc:
        assert "não reconheceu" in str(exc) or "ambíguo" in str(exc)
    else:
        raise AssertionError("O arquivo ambíguo deveria ser rejeitado")

def test_reader_consolidates_template_without_new_time_columns_as_not_verified(tmp_path):
    file = tmp_path / "antigo.xlsx"; wb = Workbook(); ws = wb.active; ws.title = "Solicitações"
    row = ["" for _ in range(22)]; row[1], row[5] = "Matrícula", "Data"; ws.append(row)
    row = ["" for _ in range(22)]; row[1], row[5] = "000123", date(2026,10,28); ws.append(row); wb.save(file)
    records, _ = read_requests(file, "1. LETIVA", "letiva", profile(), tmp_path, "exec", CancellationToken())
    assert records[0].matricula == "000123"
    assert records[0].start_time is None and "Horário de início ou término ausente ou inválido" in records[0].issues
